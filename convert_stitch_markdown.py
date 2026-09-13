from pathlib import Path
import re, json, hashlib, shutil, sys
import pdfplumber
from pptx import Presentation
from PIL import Image, ImageDraw

sys.stdout.reconfigure(encoding='utf-8')
repo=Path('D:/Преза лекций/github-glava8-kit')
dest=repo/'stitch'
dest.mkdir(exist_ok=True)
(dest/'images').mkdir(exist_ok=True)

def clean(text):
    text=(text or '').replace('\x00','').replace('\u00ad','').replace('\uf0b7','•').replace('\ue12c','•')
    return '\n'.join(line.strip() for line in text.splitlines() if line.strip())

def table_md(rows):
    if not rows: return ''
    width=max(map(len,rows))
    rows=[r+['']*(width-len(r)) for r in rows]
    def row(r): return '| '+' | '.join(clean(v).replace('|','\\|').replace('\n','<br>') for v in r)+' |'
    return '\n'.join([row(rows[0]),'| '+' | '.join(['---']*width)+' |']+[row(r) for r in rows[1:]])

results=[]
for filename, output, label in [
 ('Глава 8. Транспортный уровень(лекция).pdf','01-lecture.md','Лекция'),
 ('Глава 8. Транспортный уровень.pdf','02-handout.md','Учебный материал')]:
    parts=[f'# {label}: транспортный уровень',f'Источник: `{filename}`. Полная текстовая выгрузка по страницам. Содержание источника сохранено; неточности оригинала не исправлялись молча. Изображения без текстового слоя не заменяются выдуманным текстом. Таблицы, распознанные по структуре PDF, дополнительно приведены в Markdown.']
    counts=[]
    with pdfplumber.open(repo/'sources'/filename) as pdf:
        for i,page in enumerate(pdf.pages,1):
            t=clean(page.extract_text(x_tolerance=2,y_tolerance=3))
            lines=t.splitlines()
            if lines and lines[0]==str(i): lines=lines[1:]
            if lines and lines[-1]==str(i): lines=lines[:-1]
            t='\n'.join(lines)
            parts.extend([f'## Страница {i}',t or '> На странице нет извлекаемого текстового слоя.'])
            tables=page.extract_tables()
            for j,table in enumerate(tables,1):
                if len(table)>1 and max(map(len,table))>1:
                    parts.extend([f'### Таблица {j} на странице {i}',table_md(table)])
            counts.append(len(t))
    text='\n\n'.join(parts)+'\n'
    (dest/output).write_text(text,encoding='utf-8')
    results.append({'source':filename,'output':output,'units':len(counts),'text_characters':sum(counts),'empty_units':[i+1 for i,n in enumerate(counts) if not n]})

pres=Presentation(next((repo/'sources').glob('*.pptx')))
slides=[]
image_index=[]
seen={}
contact=[]
def shapes(items):
    for s in items:
        if s.shape_type==6:
            yield from shapes(s.shapes)
        else: yield s

for i,slide in enumerate(pres.slides,1):
    pieces=[]
    image_links=[]
    for shape in shapes(slide.shapes):
        if shape.has_text_frame:
            t=clean(shape.text)
            if t: pieces.append(t)
        if shape.has_table:
            pieces.append(table_md([[c.text for c in row.cells] for row in shape.table.rows]))
        if shape.shape_type==13:
            blob=shape.image.blob
            sha=hashlib.sha256(blob).hexdigest()
            if sha not in seen:
                name=f'slide-{i:02d}-image-{len(seen)+1:02d}.{shape.image.ext}'
                (dest/'images'/name).write_bytes(blob)
                seen[sha]=name
            name=seen[sha]
            image_links.append(f'![Иллюстрация исходного слайда {i}](images/{name})')
            image_index.append({'slide':i,'file':'images/'+name})
            contact.append((i,dest/'images'/name))
    if slide.has_notes_slide:
        notes=clean(slide.notes_slide.notes_text_frame.text)
        if notes and notes!=str(i): pieces.append('### Заметки докладчика\n\n'+notes)
    slides.append(f'## Слайд {i}\n\n'+'\n\n'.join(pieces+image_links))
(dest/'03-presentation.md').write_text('# Презентация: транспортный уровень\n\nПолная выгрузка 29 слайдов PPTX: текст, таблицы и ссылки на исходные изображения. Текст на растровых изображениях требует отдельного прочтения.\n\n'+'\n\n---\n\n'.join(slides)+'\n',encoding='utf-8')
(dest/'images/index.json').write_text(json.dumps(image_index,ensure_ascii=False,indent=2),encoding='utf-8')
results.append({'source':'PPTX','output':'03-presentation.md','units':len(slides),'images':len(seen)})
(dest/'conversion-report.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')

# Inspection sheet for raster-only teaching diagrams; not an output for Stitch.
sheet=Image.new('RGB',(1200,300*((len(contact)+2)//3)), 'white')
draw=ImageDraw.Draw(sheet)
for n,(i,p) in enumerate(contact):
    im=Image.open(p).convert('RGB'); im.thumbnail((385,265))
    x=(n%3)*400; y=(n//3)*300
    sheet.paste(im,(x,y+25)); draw.text((x+4,y+4),f'Slide {i}: {p.name}',fill='black')
sheet.save('D:/Преза лекций/stitch-image-review.jpg')
print(json.dumps(results,ensure_ascii=False,indent=2))
